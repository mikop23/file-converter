import sys, os, zipfile, io, json
from PIL import Image
from pathlib import Path
from playwright.sync_api import sync_playwright
S = Path(sys.argv[1]); IN = S/"in"; DL = S/"dl"
res = []
def ok(name, cond, extra=""):
    res.append(cond); print(("O " if cond else "X ") + name, extra)

with sync_playwright() as p:
    b = p.chromium.launch()
    ctx = b.new_context(accept_downloads=True, viewport={"width": 1280, "height": 900})
    page = ctx.new_page()
    errors = []
    page.on("pageerror", lambda e: errors.append(str(e)))
    page.on("console", lambda m: m.type == "error" and errors.append(m.text))
    page.goto("http://127.0.0.1:8765/index.html")
    page.wait_for_function("window.__fconv")
    page.screenshot(path=str(S/"s0.png"))
    page.set_input_files("#pick", [str(f) for f in sorted(IN.iterdir())])
    page.wait_for_timeout(500)
    ok("파일 8개 들어감", page.locator("#list li").count() == 8, page.locator("#count").inner_text())

    def grab(click, timeout=60000):
        with page.expect_download(timeout=timeout) as d:
            click()
        dl = d.value; path = DL/dl.suggested_filename; dl.save_as(path); return path

    # ① 묶기 (옵션은 접혀 있음 → 펼치기)
    ok("처음엔 옵션 접혀 있음", not page.is_visible("#sort"))
    page.click("details.more summary")
    page.select_option("#sort", "time")
    pdf = grab(lambda: page.click("[data-bundle=pdf]"))
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(str(pdf)); sizes = [doc[i].get_size() for i in range(len(doc))]
    ok("PDF 4장", len(doc) == 4, f"{pdf.name} {[tuple(round(x) for x in s) for s in sizes]}")
    # 시간순: 사진3(10/01) 사진2(10/02) 사진1(10/03, 회전→세로) 로고(EXIF 없음 → 파일 시간)
    portrait = [s[1] > s[0] for s in sizes]
    ok("회전 보정 (사진1이 세로)", portrait.count(True) == 1, str(portrait))
    page.select_option("#sort", "none")
    pptx = grab(lambda: page.click("[data-bundle=pptx]"))
    from pptx import Presentation
    ok("PPT 4장", len(Presentation(str(pptx)).slides) == 4, pptx.name)
    page.select_option("#perPage", "2")
    dx = grab(lambda: page.click("[data-bundle=docx]"))
    from docx import Document
    d = Document(str(dx)); n_img = len(d.inline_shapes)
    breaks = sum(1 for p in d.paragraphs for r in p.runs if 'w:br' in r._r.xml and 'type="page"' in r._r.xml)
    ok("Word 4장, 2장씩(쪽 나눔 1번)", n_img == 4 and breaks == 1, f"그림 {n_img} 쪽나눔 {breaks}")
    page.select_option("#quality", "small")
    z = grab(lambda: page.click("[data-bundle=zip]"))
    names = zipfile.ZipFile(z).namelist()
    ok("ZIP 4장 + 순서 번호", len(names) == 4 and names[0].startswith("001_"), str(names))
    ok("용량 작게 = 긴 변 1500", max(Image.open(zipfile.ZipFile(z).open(names[1])).size) == 1500)

    # ①-2 사진 형식 바꾸기 (4장 → ZIP 하나)
    page.select_option("#quality", "normal")
    pz = grab(lambda: page.click("[data-imgconv=png]"))
    zn = zipfile.ZipFile(pz).namelist()
    ok("사진 4장 → PNG ZIP", pz.name == "사진_png.zip" and len(zn) == 4 and all(n.endswith(".png") for n in zn), str(zn))
    big = Image.open(zipfile.ZipFile(pz).open("사진2.png"))
    ok("형식 바꾸기는 원본 크기 유지", big.size == (2400, 1600) and big.format == "PNG", str(big.size))

    # ② PDF 쪼개기
    sp = grab(lambda: page.click("#splitBtn"))
    ok("PDF → 이미지 2장", len(zipfile.ZipFile(sp).namelist()) == 2, sp.name)

    # ③ 문서 변환
    got = []
    page.on("download", lambda dl: got.append(dl))
    def conv(target, count=3):
        page.select_option("#target", target)
        got.clear()
        page.click("#convBtn")
        page.wait_for_function("!window.__fconv.state.busy")
        page.wait_for_timeout(300)
        paths = []
        for dl in got:
            pth = DL/dl.suggested_filename; dl.save_as(pth); paths.append(pth)
        if len(paths) != count: print("   받은 파일:", [x.name for x in paths], page.inner_text("#msg"))
        return paths
    docs = conv("docx", 2)
    md_docx = Document(str(DL/"메모.docx"))
    import re as _re
    styles = set(_re.findall(r'w:pStyle w:val="([^"]+)"', zipfile.ZipFile(DL/"메모.docx").read("word/document.xml").decode()))
    ok("md → docx (제목·목록·표)", "Heading1" in styles and len(md_docx.tables) == 1 and "ListParagraph" in styles,
       f"{[x.name for x in docs]} 표:{len(md_docx.tables)}")
    html_docx = Document(str(DL/"소개.docx"))
    ok("html → docx (script 제거)", "alert" not in " ".join(p.text for p in html_docx.paragraphs), [p.text for p in html_docx.paragraphs][:3])
    mds = conv("md", 2)
    md_text = (DL/"보고서.md").read_text()
    ok("docx → md", "# 보고서 제목" in md_text and "본문 문단" in md_text, md_text[:60].replace("\n", " | "))
    txts = conv("txt")
    ok("md → txt", "실험 메모" in (DL/"메모.txt").read_text(), (DL/"메모.txt").read_text()[:40].replace("\n", " | "))
    htmls = conv("html")
    ok("md → html", "<table>" in (DL/"메모.html").read_text())

    # ④ 말로 시키기 (규칙 모드)
    ok("AI 표시 = AI 꺼짐", page.inner_text("#aiChip") == "AI 꺼짐", page.inner_text("#aiChip"))
    page.fill("#ask", "사진 피피티로 묶어줘")
    page.click("#askBtn")
    page.wait_for_selector("#confirmDlg[open]")
    body = page.inner_text("#cfBody"); title = page.inner_text("#cfSub")
    ppt2 = grab(lambda: page.click("#confirmDlg button[value=yes]"))
    ok("글로 → PPT (AI 없이)", ppt2.suffix == ".pptx" and "문장에서 찾은" in title, body.replace("\n", " "))
    page.fill("#ask", "로고 jpg로 바꿔줘")
    page.click("#askBtn"); page.wait_for_selector("#confirmDlg[open]")
    lj = grab(lambda: page.click("#confirmDlg button[value=yes]"))
    li = Image.open(lj)
    ok("글로 → 로고 PNG를 JPG로 (투명→흰 배경)", lj.name == "로고.jpg" and li.format == "JPEG" and li.getpixel((5, 5))[0] > 150, f"{lj.name} {li.getpixel((5,5))}")
    page.fill("#ask", "스캔 글자 뽑아줘")
    page.click("#askBtn")
    page.wait_for_timeout(500)
    ok("오프라인에선 글자 추출 거절", "✗" in page.inner_text("#msg"), page.inner_text("#msg"))
    page.screenshot(path=str(S/"s1.png"), full_page=True)

    # ⑤ 온라인 AI (진짜 중계 서버 코드 + 가짜 Claude)
    page.click("#serverBtn"); page.fill("#serverUrl", "http://127.0.0.1:18788"); page.click("#serverDlg button[value=yes]")
    page.wait_for_function("document.querySelector('#aiChip').textContent === 'AI 켜짐'", timeout=8000)
    ok("설정에 주소 → AI 켜짐", True, page.inner_text("#aiChip"))
    page.fill("#ask", "사진 워드에 2장씩 넣어줘")      # 가짜 AI는 일부러 'split'으로 틀리게 답함
    page.click("#askBtn"); page.wait_for_selector("#confirmDlg[open]")
    t2 = page.inner_text("#cfSub"); b2 = page.inner_text("#cfBody")
    w2 = grab(lambda: page.click("#confirmDlg button[value=yes]"))
    ok("AI 틀린 답 보정 → Word", "AI가 이해한" in t2 and w2.suffix == ".docx", b2.replace("\n", " "))
    page.fill("#ask", "스캔 글자 뽑아줘")
    page.click("#askBtn"); page.wait_for_selector("#confirmDlg[open]")
    page.click("#confirmDlg button[value=yes]")            # 작업 확인
    page.wait_for_selector("#confirmDlg[open]")
    consent = page.inner_text("#cfTitle")
    ocr = grab(lambda: page.click("#confirmDlg button[value=yes]"))   # 온라인 전송 안내
    od = Document(str(ocr))
    ok("글자 뽑기 → Word (표 포함, 2쪽)", consent == "사진을 AI로 보낼게요" and len(od.tables) == 2, f"{ocr.name} 표 {len(od.tables)}")
    page.screenshot(path=str(S/"s2.png"), full_page=True)

    # 폰 화면 + 다크 모드
    m = b.new_context(viewport={"width": 390, "height": 844}, color_scheme="dark", device_scale_factor=2)
    mp = m.new_page(); mp.goto("http://127.0.0.1:8765/index.html"); mp.wait_for_function("window.__fconv")
    mp.set_input_files("#pick", [str(IN/"사진1.jpg"), str(IN/"메모.md")]); mp.wait_for_timeout(400)
    sw = mp.evaluate("document.documentElement.scrollWidth")
    ok("폰 화면 가로 스크롤 없음", sw <= 390, f"scrollWidth {sw}")
    ok("파일 있으면 안내 문구 숨김", not mp.is_visible("#empty"))
    ok("PDF 없으면 'PDF를 사진으로' 카드에 이유 표시", mp.is_visible("#cardSplit .why") and "off" in mp.get_attribute("#cardSplit", "class"))
    mp.screenshot(path=str(S/"s3_mobile_dark.png"), full_page=True)
    errors = [e for e in errors if "ERR_TUNNEL" not in e]   # 테스트 환경에서 웹폰트 CDN 막힘 (사이트 문제 아님)
    ok("브라우저 오류 없음", not errors, str(errors[:3]))
    b.close()
print(f"\n{sum(res)}/{len(res)} 통과")
