#!/usr/bin/env python3
"""
fconv - 문서 변환 + 이미지 묶기/쪼개기 도구 (Termux / PC 공용)

사용 예:
  python fconv.py convert 보고서.md docx          # 문서 변환 (pandoc)
  python fconv.py convert 과제.pdf docx            # PDF → Word (pdf2docx)
  python fconv.py bundle 사진*.jpg -o 과제.pdf     # 이미지 여러 장 → PDF
  python fconv.py bundle *.jpg -o 앨범.pptx --sort time
  python fconv.py bundle *.png -o 모음.docx --per-page 2
  python fconv.py bundle *.jpg -o 묶음.zip
  python fconv.py split 스캔본.pdf -f jpg          # PDF → 이미지 여러 장

이미지 묶기 옵션:
  --sort name|time   파일명순(기본, 1,2,10 자연정렬) / 찍은 시간순(EXIF)
  --quality 1~95     JPEG 압축 품질 (기본 80, 낮을수록 용량↓)
  --max-size 픽셀    긴 변 최대 크기 (기본 2000, 0이면 원본 유지)
  회전 자동 보정은 항상 켜져 있음 (폰 사진 눕는 문제 해결)
"""
import argparse, io, os, re, shutil, subprocess, sys, tempfile, zipfile
from datetime import datetime
from pathlib import Path

try:
    from PIL import Image, ImageOps
except ImportError:
    sys.exit("Pillow 필요: pip install pillow")

IMG_EXT = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".heic"}
PANDOC_FMT = {"md": "markdown", "markdown": "markdown", "docx": "docx", "html": "html",
              "htm": "html", "txt": "plain", "odt": "odt", "epub": "epub", "rtf": "rtf",
              "tex": "latex", "rst": "rst", "ipynb": "ipynb"}


# ───────── 공통 ─────────
APP_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))

def find_tool(name):
    """프로그램 옆 bin 폴더(배포용) → PATH 순서로 외부 도구 찾기"""
    for c in (APP_DIR / "bin" / (name + ".exe"), APP_DIR / "bin" / name):
        if c.is_file():
            return str(c)
    return shutil.which(name)

def find_office():
    return find_tool("soffice") or shutil.which("libreoffice") or next(
        (str(c) for c in [Path(os.environ.get(v, "")) / "LibreOffice/program/soffice.exe"
                          for v in ("ProgramFiles", "ProgramFiles(x86)")] if c.is_file()), None)

def unique_path(p):
    """같은 이름 있으면 '이름 (1).pdf' 식으로 → 덮어쓰기 방지"""
    p = Path(p)
    if not p.exists():
        return p
    i = 1
    while True:
        c = p.with_name(f"{p.stem} ({i}){p.suffix}")
        if not c.exists():
            return c
        i += 1

NO_WINDOW = {"creationflags": 0x08000000} if os.name == "nt" else {}  # 윈도우에서 검은 창 안 뜨게


# ───────── 공통 이미지 처리 (덤 기능) ─────────
def natural_key(p):
    """'img2' < 'img10' 처럼 사람이 생각하는 순서로 정렬"""
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", Path(p).name)]

def taken_time(p):
    """EXIF 촬영시간, 없으면 파일 수정시간"""
    try:
        exif = Image.open(p).getexif()
        raw = exif.get_ifd(0x8769).get(36867) or exif.get(306)
        if raw:
            return datetime.strptime(raw, "%Y:%m:%d %H:%M:%S").timestamp()
    except Exception:
        pass
    return os.path.getmtime(p)

def sort_images(paths, mode):
    if mode == "none":                                 # 사용자가 정한 순서 그대로
        return list(paths)
    return sorted(paths, key=taken_time if mode == "time" else natural_key)

def load_image(p, max_size):
    """회전 보정 + RGB 변환 + 크기 축소"""
    if Path(p).suffix.lower() == ".heic":
        try:
            import pillow_heif; pillow_heif.register_heif_opener()
        except ImportError:
            sys.exit("HEIC 사진은 pip install pillow-heif 필요")
    im = ImageOps.exif_transpose(Image.open(p))       # ← 회전 자동 보정
    if im.mode in ("RGBA", "LA", "P"):
        bg = Image.new("RGB", im.size, "white")
        im = im.convert("RGBA"); bg.paste(im, mask=im.split()[-1]); im = bg
    else:
        im = im.convert("RGB")
    if max_size and max(im.size) > max_size:
        im.thumbnail((max_size, max_size))            # ← 용량 줄이기 (해상도)
    return im

def to_jpeg_bytes(im, quality):
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=quality, optimize=True)  # ← 용량 줄이기 (압축)
    buf.seek(0)
    return buf


# ───────── ① 문서 변환 ─────────
def convert(src, target, out=None):
    src = Path(src)
    target = target.lower().lstrip(".")
    out = Path(out) if out else unique_path(src.with_suffix("." + target))
    s = src.suffix.lower().lstrip(".")
    if s == target:
        sys.exit(f"{src.name}: 이미 {target} 형식")

    if s == "pdf" and target == "docx":
        try:
            from pdf2docx import Converter
        except ImportError:
            sys.exit("PDF→Word는 pip install pdf2docx 필요")
        cv = Converter(str(src)); cv.convert(str(out)); cv.close()
    elif target == "pdf" and s in {"docx", "doc", "pptx", "ppt", "xlsx", "xls", "odt", "hwp"}:
        office = find_office()
        if not office:
            sys.exit("Office→PDF는 LibreOffice 설치 필요")
        with tempfile.TemporaryDirectory() as td:           # 임시 폴더에서 만들고 옮김
            subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", td, str(src)],
                           check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, **NO_WINDOW)
            made = Path(td) / (src.stem + ".pdf")
            if not made.exists():
                sys.exit("LibreOffice 변환 실패")
            shutil.move(str(made), str(out))
    elif "." + s in IMG_EXT and target in {"pdf", "jpg", "jpeg", "png", "webp"}:
        return bundle([src], out, quiet=True)
    elif s in PANDOC_FMT and target == "pdf":
        # md/html/txt → pdf : pandoc으로 docx 만든 뒤 LibreOffice로 pdf (LaTeX 설치 불필요)
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td) / (src.stem + ".docx")
            convert(src, "docx", tmp)
            convert(tmp, "pdf", out)
        return out
    elif s in PANDOC_FMT and target in PANDOC_FMT:
        pandoc = find_tool("pandoc")
        if not pandoc:
            sys.exit("pandoc 설치 필요")
        r = subprocess.run([pandoc, str(src), "-f", PANDOC_FMT[s], "-t", PANDOC_FMT[target],
                            "-o", str(out)], capture_output=True, text=True, **NO_WINDOW)
        if r.returncode:
            sys.exit("pandoc 오류: " + r.stderr.strip()[:200])
    else:
        sys.exit(f"{s} → {target} 변환은 아직 지원 안 함")
    print(f"완료: {out}")
    return out


# ───────── ② 이미지 묶기 ─────────
def bundle(paths, out, sort="name", quality=80, max_size=2000, per_page=1, quiet=False):
    import glob   # 윈도우 cmd는 *.jpg를 안 펼쳐줘서 직접 펼침
    paths = [g for p in paths for g in (glob.glob(str(p)) if any(c in str(p) for c in "*?[") else [p])]
    paths = [p for p in paths if Path(p).suffix.lower() in IMG_EXT]
    if not paths:
        sys.exit("이미지 파일이 없음")
    paths = sort_images(paths, sort)
    out = unique_path(out); kind = out.suffix.lower()

    if kind == ".pdf":
        imgs = [Image.open(to_jpeg_bytes(load_image(p, max_size), quality)) for p in paths]
        imgs[0].save(out, "PDF", save_all=True, append_images=imgs[1:], resolution=150)

    elif kind == ".docx":
        from docx import Document
        from docx.shared import Cm
        doc = Document()
        sec = doc.sections[0]
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, m, Cm(1.5))
        usable_w = sec.page_width - sec.left_margin - sec.right_margin
        usable_h = (sec.page_height - sec.top_margin - sec.bottom_margin) / per_page * 0.95
        for i, p in enumerate(paths):
            im = load_image(p, max_size)
            ratio = im.width / im.height
            w = min(usable_w, int(usable_h * ratio))
            doc.add_picture(to_jpeg_bytes(im, quality), width=w)
            if (i + 1) % per_page == 0 and i + 1 < len(paths):
                doc.add_page_break()
        doc.save(out)

    elif kind == ".pptx":
        from pptx import Presentation
        from pptx.util import Emu
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
        W, H = prs.slide_width, prs.slide_height
        for p in paths:
            im = load_image(p, max_size)
            s = prs.slides.add_slide(prs.slide_layouts[6])
            scale = min(W / im.width, H / im.height)
            w, h = int(im.width * scale), int(im.height * scale)
            s.shapes.add_picture(to_jpeg_bytes(im, quality), (W - w) // 2, (H - h) // 2, w, h)
        prs.save(out)

    elif kind == ".zip":
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for i, p in enumerate(paths, 1):
                z.writestr(f"{i:03d}_{Path(p).stem}.jpg",
                           to_jpeg_bytes(load_image(p, max_size), quality).read())

    elif kind in {".jpg", ".jpeg", ".png", ".webp"} and len(paths) == 1:
        im = load_image(paths[0], max_size)
        im.save(out, quality=quality) if kind != ".png" else im.save(out, optimize=True)

    else:
        sys.exit(f"{kind} 로 묶기는 지원 안 함 (pdf/docx/pptx/zip)")

    size = out.stat().st_size / 1024 / 1024
    if not quiet:
        print(f"완료: {out}  ({len(paths)}장, {size:.2f}MB)")
    return out


# ───────── ③ PDF → 이미지 쪼개기 ─────────
def split(pdf, fmt="jpg", dpi=150, outdir=None):
    pdf = Path(pdf)
    outdir = Path(outdir) if outdir else unique_path(pdf.with_name(pdf.stem + "_pages"))
    outdir.mkdir(parents=True, exist_ok=True)
    ext = "jpg" if fmt in ("jpg", "jpeg") else "png"
    try:
        import pypdfium2 as pdfium                         # pip만으로 설치됨 (배포에 유리)
        doc = pdfium.PdfDocument(str(pdf))
        for i, page in enumerate(doc, 1):
            im = page.render(scale=dpi / 72).to_pil().convert("RGB")
            im.save(outdir / f"page-{i:03d}.{ext}", quality=90)
        n = len(doc); doc.close()
    except ImportError:
        tool = find_tool("pdftoppm")
        if not tool:
            sys.exit("PDF 쪼개기는 pip install pypdfium2 필요")
        subprocess.run([tool, "-jpeg" if ext == "jpg" else "-png", "-r", str(dpi), str(pdf),
                        str(outdir / "page")], check=True, **NO_WINDOW)
        n = len(list(outdir.iterdir()))
    print(f"완료: {outdir}/ 에 {n}장")
    return outdir


# ───────── CLI ─────────
def main():
    ap = argparse.ArgumentParser(description="문서 변환 + 이미지 묶기 도구")
    sub = ap.add_subparsers(dest="cmd", required=True)

    c = sub.add_parser("convert", help="문서 형식 변환")
    c.add_argument("src"); c.add_argument("target", help="docx, pdf, md, html, txt ...")
    c.add_argument("-o", "--out")

    b = sub.add_parser("bundle", help="이미지 여러 장 → pdf/docx/pptx/zip")
    b.add_argument("images", nargs="+"); b.add_argument("-o", "--out", required=True)
    b.add_argument("--sort", choices=["name", "time", "none"], default="name")
    b.add_argument("--quality", type=int, default=80)
    b.add_argument("--max-size", type=int, default=2000)
    b.add_argument("--per-page", type=int, default=1, help="Word 한 페이지당 이미지 수")

    s = sub.add_parser("split", help="PDF → 이미지 여러 장")
    s.add_argument("pdf"); s.add_argument("-f", "--format", default="jpg", choices=["jpg", "png"])
    s.add_argument("--dpi", type=int, default=150); s.add_argument("-o", "--outdir")

    a = ap.parse_args()
    if a.cmd == "convert": convert(a.src, a.target, a.out)
    elif a.cmd == "bundle": bundle(a.images, a.out, a.sort, a.quality, a.max_size, a.per_page)
    elif a.cmd == "split": split(a.pdf, a.format, a.dpi, a.outdir)

if __name__ == "__main__":
    main()
